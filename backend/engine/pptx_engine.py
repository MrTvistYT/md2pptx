from logging import config

from pptx import Presentation
from io import BytesIO

from pptx.util import Inches

from engine.exceptions import UndefinedFormat

class Slide:
    def add_to_presentation(self, presentation: Presentation):
        pass

class TitleSlide(Slide):
    def __init__(self, title, subtitle):
        self.title = title
        self.subtitle = subtitle

    def add_to_presentation(self, presentation: Presentation):
        title_slide_layout = presentation.slide_layouts[0]

        slide = presentation.slides.add_slide(title_slide_layout)

        title = slide.shapes.title
        title.text = self.title

        subtitle = slide.placeholders[1]
        subtitle.text = self.subtitle

class BulletContent(Slide):
    def __init__(self, subtitle: str, content: str):
        self.subtitle = subtitle
        self.content = content

    def add_to_presentation(self, presentation: Presentation):
        bullet_slide_layout = presentation.slide_layouts[1]

        slide = presentation.slides.add_slide(bullet_slide_layout)

        title_shape = slide.shapes.title
        title_shape.text = self.subtitle

        body_shape = slide.shapes.placeholders[1]
        tf = body_shape.text_frame
        p = tf.add_paragraph()
        p.text = self.content

class ContentReader:
    def __init__(self, pptx_content: str):
        self.pptx_content = pptx_content

    def slides(self) -> list[str]:
        return self.pptx_content.split("---\n")

    def define_slide_type(self, slide: str) -> Slide:
        textboxes = slide.split("\n")

        titles = []
        subtitles = []
        texts = []

        for textbox in textboxes:
            if not textbox:
                continue

            if textbox.startswith("##"):
                subtitles.append(textbox[2:])
            elif textbox.startswith("#"):
                titles.append(textbox[1:])
            else:
                texts.append(textbox)


        if titles and subtitles:
            return TitleSlide(titles[0], subtitles[0])
        elif subtitles and texts:
            return BulletContent(subtitles[0], texts[0])

        raise UndefinedFormat

class PresentationGenerator:
    def __init__(self, pptx_content: str):
        self.presentation = Presentation()
        self.presentation.slide_width = Inches(13.33)
        self.presentation.slide_height = Inches(7.5)

        self.reader = ContentReader(pptx_content)

    def unparse(self):
        for slide in self.reader.slides():
            slide_type = self.reader.define_slide_type(slide)
            slide_type.add_to_presentation(self.presentation)

    def export(self) -> BytesIO:
        pptx_stream = BytesIO()
        self.presentation.save(pptx_stream)
        pptx_stream.seek(0)
        return pptx_stream
